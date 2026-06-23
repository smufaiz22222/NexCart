from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_PATH = PROJECT_ROOT / "docs" / "NexCart_Security_Remediation_Update.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color="000000", size=10.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)


def add_body(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)


def main():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("NexCart Security Remediation Update")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor.from_string("1F3A5F")
    title_run.bold = True

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle_run = subtitle.add_run(
        f"Prepared on {date(2026, 6, 23).strftime('%B %d, %Y')} for the latest backlog-driven fixes."
    )
    subtitle_run.font.name = "Calibri"
    subtitle_run.font.size = Pt(10.5)
    subtitle_run.font.color.rgb = RGBColor.from_string("5B6572")

    add_body(
        document,
        "This memo captures the remediation work completed for the latest security and performance findings raised during the code audit. It focuses on access control for the AI advisor service, browser token handling, environment-secret exposure guidance, CORS hardening, and the remaining Khatta write bottleneck.",
    )

    summary_heading = document.add_paragraph()
    summary_heading.paragraph_format.space_before = Pt(10)
    summary_heading.paragraph_format.space_after = Pt(6)
    summary_run = summary_heading.add_run("Remediation Summary")
    summary_run.bold = True
    summary_run.font.name = "Calibri"
    summary_run.font.size = Pt(15)
    summary_run.font.color.rgb = RGBColor.from_string("1F3A5F")

    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.45), Inches(1.1), Inches(1.95), Inches(2.0)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width

    headers = ["Finding", "Priority", "Change", "Result"]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "EAF0F6")
        set_cell_text(cell, header, bold=True, color="1F3A5F")

    summary_rows = [
        (
            "AI advisor service exposed publicly",
            "Critical",
            "Added JWT validation, role-gated routes, and user-scoped session history.",
            "Anonymous chat/history/ingest access is blocked; ingest is admin-only.",
        ),
        (
            "Browser tokens stored in localStorage",
            "High",
            "Moved session auth to httpOnly cookie flow and credentialed API requests.",
            "Frontend JavaScript no longer reads or persists bearer tokens.",
        ),
        (
            "Plaintext local secrets in workspace",
            "High",
            "Documented that the local .env values must be rotated if real.",
            "Risk is operationally contained only after secret rotation by maintainers.",
        ),
        (
            "Unsafe AI service CORS fallback",
            "Medium",
            "Removed wildcard fallback and require explicit AI_CORS_ORIGINS values.",
            "Credentialed cross-origin calls are now restricted to configured origins.",
        ),
        (
            "Khatta transactional write loop",
            "Medium",
            "Replaced per-row inserts with batched createManyAndReturn in one transaction.",
            "Invoice saves scale better and avoid N+1 write behavior.",
        ),
    ]

    for finding, priority, change, result in summary_rows:
        row = table.add_row()
        values = [finding, priority, change, result]
        for idx, value in enumerate(values):
            set_cell_text(row.cells[idx], value)

    sections = [
        (
            "1. AI Advisor Service Access Control",
            [
                "The FastAPI advisor service now validates the shared NexCart JWT before serving business-advisor endpoints.",
                "The /chat and /history routes require an authenticated WHOLESALER or SUPER_ADMIN user, while /ingest is limited to SUPER_ADMIN users only.",
                "Session history is scoped as userId:sessionId, which prevents one logged-in user from reading another user's transcript by guessing a session identifier.",
                "The frontend AI client was updated to send credentialed requests instead of calling the advisor service anonymously.",
            ],
        ),
        (
            "2. Browser Session Hardening",
            [
                "Authentication tokens are no longer written to localStorage or persisted in the Zustand auth store.",
                "Login now sets an httpOnly cookie on the backend, and logout clears that cookie while preserving token revocation behavior.",
                "The frontend now restores session state from /auth/profile and uses withCredentials for both the main API client and the AI advisor client.",
                "A non-sensitive session hint is used where the UI only needs to know whether a customer session exists, reducing exposure to token theft through future XSS.",
            ],
        ),
        (
            "3. Secret Exposure Response",
            [
                "The local .env file remains untracked, so this was treated as a workspace exposure rather than confirmed source-control leakage.",
                "No new .env.example files were retained because that was explicitly not wanted.",
                "If the database, JWT, Gemini, Razorpay, or Brevo values in the local .env are real, they should still be treated as compromised and rotated outside the codebase.",
            ],
        ),
        (
            "4. CORS Hardening For The AI Service",
            [
                "The AI service no longer falls back to allow_origins=[\"*\"] when credentials are enabled.",
                "Startup now fails fast if AI_CORS_ORIGINS is blank, which prevents accidental deployment of a broadly callable credentialed advisor origin policy.",
            ],
        ),
        (
            "5. Khatta Save Path Performance Fix",
            [
                "The save flow still validates extracted customer rows inside a single Prisma transaction, but it no longer issues one awaited insert per ledger row.",
                "Resolved rows are collected into one payload and inserted with createManyAndReturn, removing the prior N+1 transactional write loop.",
            ],
        ),
    ]

    for heading_text, bullets in sections:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(4)
        run = heading.add_run(heading_text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor.from_string("1F3A5F")

        for bullet in bullets:
            add_bullet(document, bullet)

    verification_heading = document.add_paragraph()
    verification_heading.paragraph_format.space_before = Pt(12)
    verification_heading.paragraph_format.space_after = Pt(4)
    verification_run = verification_heading.add_run("Verification Notes")
    verification_run.bold = True
    verification_run.font.name = "Calibri"
    verification_run.font.size = Pt(13)
    verification_run.font.color.rgb = RGBColor.from_string("1F3A5F")

    verification_items = [
        "Backend auth controller tests were run with node --test src/controllers/authController.test.js.",
        "The Khatta controller syntax was checked with node --check src/controllers/khattaController.js.",
        "Frontend lint and AI-service syntax checks had already been run during remediation work.",
        "Visual DOCX render QA could not be completed in this environment because LibreOffice/soffice is not installed.",
    ]
    for item in verification_items:
        add_bullet(document, item)

    footer_section = document.sections[-1]
    footer_section.start_type = WD_SECTION.CONTINUOUS
    footer = footer_section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("NexCart internal remediation record")
    footer_run.font.name = "Calibri"
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor.from_string("6B7280")

    document.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
